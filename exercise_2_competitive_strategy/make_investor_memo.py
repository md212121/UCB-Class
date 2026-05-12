import matplotlib
matplotlib.use("Agg")
from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

doc = Document()

# ── Margins: tight to fit two pages ─────────────────────────────────────────
section = doc.sections[0]
section.top_margin    = Inches(0.75)
section.bottom_margin = Inches(0.75)
section.left_margin   = Inches(1.0)
section.right_margin  = Inches(1.0)

# ── Helpers ──────────────────────────────────────────────────────────────────
def rule(space_before=4, space_after=4):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after  = Pt(space_after)
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bot = OxmlElement('w:bottom')
    bot.set(qn('w:val'), 'single'); bot.set(qn('w:sz'), '4')
    bot.set(qn('w:space'), '1');    bot.set(qn('w:color'), 'CCCCCC')
    pBdr.append(bot); pPr.append(pBdr)

def h1(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(9)
    p.paragraph_format.space_after  = Pt(3)
    r = p.add_run(text.upper())
    r.bold = True; r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(0x44, 0x44, 0xCC)
    return p

def body(text, size=10.5, space_before=2, space_after=5, bold=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after  = Pt(space_after)
    r = p.add_run(text)
    r.font.size = Pt(size); r.bold = bold
    return p

def bullet(text, bold_prefix=None, size=10.5):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after  = Pt(3)
    if bold_prefix:
        r = p.add_run(bold_prefix); r.bold = True; r.font.size = Pt(size)
    r2 = p.add_run(text); r2.font.size = Pt(size)
    return p

def set_col_width(cell, width_inches):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcW = OxmlElement('w:tcW')
    tcW.set(qn('w:w'), str(int(width_inches * 1440)))
    tcW.set(qn('w:type'), 'dxa')
    tcPr.append(tcW)

def shade_cell(cell, hex_color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    tcPr.append(shd)

# ── HEADER ───────────────────────────────────────────────────────────────────
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(1)
r = p.add_run("MERIDIAN TECHNOLOGIES  ·  INVESTOR DAY  ·  MARCH 11, 2026")
r.font.size = Pt(8); r.font.color.rgb = RGBColor(0x88, 0x88, 0x88); r.bold = True

p2 = doc.add_paragraph()
p2.paragraph_format.space_after = Pt(1)
r2 = p2.add_run("CONFIDENTIAL — FOR INTERNAL USE PRIOR TO PUBLIC DISCLOSURE")
r2.font.size = Pt(8); r2.font.color.rgb = RGBColor(0xCC, 0x44, 0x44); r2.bold = True

rule(space_before=4, space_after=6)

# Title
t = doc.add_paragraph()
t.paragraph_format.space_after = Pt(4)
t.alignment = WD_ALIGN_PARAGRAPH.LEFT
rt = t.add_run("Investor Day Positioning Memo: Declaring Meridian's Strategic Identity")
rt.bold = True; rt.font.size = Pt(14)

# ── EXEC SUMMARY ─────────────────────────────────────────────────────────────
h1("Executive Summary")
body(
    "Meridian enters Investor Day at an inflection point. After eleven years as a project "
    "management platform, the company has the assets — enterprise governance infrastructure, "
    "the Helio agent framework, FedRAMP/HIPAA compliance, and deep regulated-industry "
    "relationships — to stake a credible claim to a position no competitor currently occupies: "
    "the enterprise-grade agentic work platform for regulated industries. This memo recommends "
    "that Meridian declare itself exactly that on March 11, set three-year financial targets "
    "consistent with that identity, and make three specific commitments to investors that create "
    "accountability for the transition. The alternative — remaining a 'PM platform with AI "
    "features' — cedes the agentic narrative to Atlassian and Asana, strands Meridian in a "
    "pricing position mid-market customers are already rejecting, and gives the Helio team "
    "insufficient reason to stay.",
    space_after=4
)

rule()

# ── POSITION ─────────────────────────────────────────────────────────────────
h1("The Position")
body(
    "Meridian is the agentic work platform built for regulated enterprises.",
    bold=True, size=12, space_after=4
)
body(
    "Projects are how we started. Agents are how work gets done in the next decade. "
    "Meridian combines the only enterprise-grade agent governance infrastructure in the "
    "category — audit trails, role-based agent permissions, model selection, data residency "
    "— with the project management depth that makes agents actionable in the most complex, "
    "compliance-sensitive organizations in the world. We do not compete with the model labs. "
    "We are the platform on which regulated enterprises run them safely.",
    space_after=4
)

rule()

# ── WHY ──────────────────────────────────────────────────────────────────────
h1("Why This Position, Why Now")
bullet(
    "Enterprise customers are already asking for it. In 18 of 22 enterprise advisory board "
    "sessions (Dec 2025–Feb 2026), customers' top AI request was governance: audit trails, "
    "role-based agent permissions, model selection, and data residency. Not better PM features "
    "— agent governance. We have this infrastructure. No direct competitor does.",
    bold_prefix="Customer demand is ahead of our positioning.  "
)
bullet(
    "Asana and Monday have raced to bundle AI into standard tiers, compressing margins and "
    "crowding the 'agentic bundled' quadrant. Smartsheet is decelerating in 'conservative PM.' "
    "Atlassian owns 'agentic premium' for dev/IT teams but has a thin story outside engineering "
    "organizations. The regulated-industry agentic platform position — financial services, "
    "pharma, life sciences, federal — is unoccupied.",
    bold_prefix="The competitive white space is real and specific.  "
)
bullet(
    "The Helio acquisition gave us a working agent orchestration framework that would have "
    "taken 15 months to build internally. The 2026 roadmap — agent governance suite (Q1), "
    "agent builder (Q2), workflow marketplace (Q3), pricing model refresh (Q3) — is already "
    "the Option B roadmap. We are not pivoting to this position; we are naming what we are building.",
    bold_prefix="We are already building this.  "
)
bullet(
    "With $506M in liquidity, $67.5M in 2025 FCF, and $250M identified M&A capacity, "
    "Meridian can fund the engineering hiring (25 applied AI engineers, 8 agent PMs), "
    "the two additional AI-native acquisitions in early diligence, and the mid-market "
    "pricing adjustment — without taking on debt or compromising the buyback.",
    bold_prefix="We can afford it.  "
)

rule()

# ── COMPETITIVE LANDSCAPE TABLE ──────────────────────────────────────────────
h1("Competitive Landscape")

p_img = doc.add_paragraph()
p_img.paragraph_format.space_before = Pt(2)
p_img.paragraph_format.space_after  = Pt(4)
try:
    from docx.shared import Inches
    run = p_img.add_run()
    run.add_picture(
        "/home/user/UCB-Class/exercise_2_competitive_strategy/positioning_matrix.png",
        width=Inches(5.5)
    )
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(5)
    cr = cap.add_run(
        "Figure 1: Competitive positioning — X: PM-centric → Agentic-first; "
        "Y: Bundled pricing → Premium/add-on. Source: public competitor statements, Feb 2026."
    )
    cr.italic = True; cr.font.size = Pt(8)
    cr.font.color.rgb = RGBColor(0x77, 0x77, 0x77)
except Exception as e:
    p_img.add_run(f"[Chart: positioning_matrix.png — {e}]").font.size = Pt(9)

# Companion table
cols  = ["Competitor", "AI Posture", "Pricing Model", "Meridian's Edge"]
widths = [0.95, 1.55, 1.3, 2.15]
rows_data = [
    ["Asana",
     "\"Agent OS\" — AI bundled in Advanced+",
     "Bundled; no consumption",
     "Governance depth Asana lacks; regulated-industry compliance Asana cannot match"],
    ["Monday.com",
     "\"Work OS\" — agents bundled in Pro+",
     "Bundled; no consumption",
     "Enterprise governance, vertical depth; Monday has no FedRAMP, limited HIPAA"],
    ["Smartsheet",
     "\"Trust-first AI\" — AI Compliance Pack add-on",
     "Bundled base; premium compliance add-on",
     "Agentic ambition; Smartsheet explicitly rejects agentic narrative"],
    ["Atlassian (Rovo)",
     "\"Agentic enterprise\" — Rovo agent builder",
     "Per-seat + consumption",
     "Regulated non-dev verticals; Atlassian strong in engineering orgs only"],
]

tbl = doc.add_table(rows=1 + len(rows_data), cols=len(cols))
tbl.style = 'Table Grid'

# Header row
hdr = tbl.rows[0]
for i, (col, w) in enumerate(zip(cols, widths)):
    cell = hdr.cells[i]
    set_col_width(cell, w)
    shade_cell(cell, "1E3A5F")
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after  = Pt(3)
    r = p.add_run(col)
    r.bold = True; r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

# Data rows
fill_colors = ["F8F9FA", "FFFFFF", "F8F9FA", "FFFFFF"]
for ri, (row_data, fill) in enumerate(zip(rows_data, fill_colors)):
    row = tbl.rows[ri + 1]
    for ci, (val, w) in enumerate(zip(row_data, widths)):
        cell = row.cells[ci]
        set_col_width(cell, w)
        shade_cell(cell, fill)
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after  = Pt(2)
        r = p.add_run(val)
        r.font.size = Pt(8.5)
        if ci == 0:
            r.bold = True

doc.add_paragraph().paragraph_format.space_after = Pt(2)
rule(space_before=2, space_after=4)

# ── THREE RISKS ───────────────────────────────────────────────────────────────
h1("Three Risks We Are Managing Explicitly")
bullet(
    "Enterprise agentic adoption may be slower than category narrative suggests, creating "
    "an air gap between today's $3.5M Copilot ARR and the consumption revenue needed to "
    "justify the positioning. Mitigation: 2026 guidance does not assume agentic consumption "
    "as a dominant contributor; the PM motion funds the transition. We will provide a "
    "specific Copilot attach-rate target (>40% enterprise renewals) as the leading indicator "
    "investors should watch.",
    bold_prefix="Revenue air gap during transition.  "
)
bullet(
    "Microsoft, Anthropic, or OpenAI enterprise agreements may make the agentic platform "
    "decision above Meridian's sales level, bypassing the buying motion entirely. Mitigation: "
    "model-neutral positioning and BYO-model support are non-negotiable product commitments; "
    "we do not tie Meridian's value to any single lab.",
    bold_prefix="Platform bypass by hyperscalers or model labs.  "
)
bullet(
    "The Helio team — 28 engineers whose work is the technical foundation of Option B — "
    "faces a compensation cliff in 2026. If they leave, the agent builder roadmap is exposed "
    "by 15 months. Mitigation: senior engineering comp refresh in Q1 2026; Helio retention "
    "packages extended; agent builder roadmap accelerated to reduce single-team dependency.",
    bold_prefix="Helio retention risk undermines the roadmap.  "
)

rule()

# ── THREE COMMITMENTS ─────────────────────────────────────────────────────────
h1("Three Commitments to Investors")
bullet(
    "We will report Copilot attach rate on enterprise renewals and in-product agentic "
    "action usage every quarter beginning Q1 2026. These are the leading indicators of "
    "whether the agentic platform is earning its positioning — and we will not hide behind "
    "aggregate ARR if they are moving in the wrong direction.",
    bold_prefix="Transparency on AI metrics, starting Q1 2026.  "
)
bullet(
    "We will announce revised mid-market Copilot pricing — moving from a standalone "
    "$40/seat add-on to a mid-market bundle — by Q2 2026. We will simultaneously "
    "introduce consumption pricing for enterprise agent actions in H2 2026. "
    "The pricing model will match the platform identity.",
    bold_prefix="Pricing model aligned to positioning by H2 2026.  "
)
bullet(
    "We will return to a three-year plan with specific financial targets: $700M–$900M "
    "revenue by 2028 (range reflects agentic ramp variance), operating margin of 16–20%, "
    "and enterprise ARR above $350M. The wide range is honest. The floor assumes the PM "
    "motion holds; the ceiling assumes agentic consumption ramps on plan. "
    "We will update the range annually.",
    bold_prefix="Three-year financial targets with a floor and a ceiling.  "
)

rule()

# Footer
f = doc.add_paragraph()
f.paragraph_format.space_before = Pt(4)
rf = f.add_run(
    "Prepared by: Office of the CEO  ·  February 2026  ·  Confidential  ·  "
    "Source data: meridian_internal_brief.md, meridian_ai_strategy_options.md, "
    "meridian_recent_customer_feedback.md, meridian_financials_summary.csv, "
    "competitors_cached/ (Feb 2026 snapshots)"
)
rf.font.size = Pt(7.5); rf.italic = True
rf.font.color.rgb = RGBColor(0x99, 0x99, 0x99)

doc.save("/home/user/UCB-Class/exercise_2_competitive_strategy/investor_day_positioning_memo.docx")
print("Saved.")
